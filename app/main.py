"""Brandproof web app (Starlette + Jinja2, server-rendered)."""
from __future__ import annotations

import base64
import csv
import io
import json
import logging
import os
import re
from contextlib import asynccontextmanager
from datetime import date

from markupsafe import Markup, escape
from starlette.applications import Starlette
from starlette.concurrency import run_in_threadpool
from starlette.exceptions import HTTPException
from starlette.requests import Request
from starlette.responses import JSONResponse, PlainTextResponse, RedirectResponse, Response
from starlette.routing import Mount, Route
from starlette.staticfiles import StaticFiles
from starlette.templating import Jinja2Templates

from .ai import AIGateway
from .checks.image import normalise_image
from .checks.pipeline import run_checks
from .db import Database, now_iso
from .kit import DIMENSION_LABELS, DIMENSIONS, load_kit
from .seed import SAMPLE_EMAIL, SAMPLE_POST, seed_kit
from .textutil import find_span, tidy_spaces, word_diff_html

logging.basicConfig(level=os.environ.get("LOG_LEVEL", "INFO"))
log = logging.getLogger("brandproof")

HERE = os.path.dirname(os.path.abspath(__file__))
MAX_TEXT = 8000
MAX_UPLOAD = 6 * 1024 * 1024

ROLES = {
    "creator": {"name": "Priya", "label": "Creator", "about": "Writes drafts, fixes findings, submits for review"},
    "reviewer": {"name": "Meera", "label": "Reviewer · Compliance", "about": "Approves drafts or requests changes"},
    "brand": {"name": "Arjun", "label": "Brand manager", "about": "Owns the Brand Kit; can also review"},
}
STATUS = {
    "checked": ("Checked", "Fix the findings, then submit"),
    "in_review": ("In review", "Waiting for a reviewer"),
    "changes_requested": ("Changes requested", "The reviewer asked for changes"),
    "approved": ("Approved", "Version locked with an approval record"),
}
CHANNELS = ["Instagram", "LinkedIn", "X / Twitter", "Email", "Blog", "Website", "Ad", "Brochure", "Other"]
EDITABLE = {"checked", "changes_requested"}

db: Database | None = None
ai: AIGateway | None = None
templates = Jinja2Templates(directory=os.path.join(HERE, "templates"))

try:
    from zoneinfo import ZoneInfo
    DISPLAY_TZ = ZoneInfo(os.environ.get("DISPLAY_TZ", "Asia/Kolkata"))
except Exception:  # no tz database available
    from datetime import timedelta, timezone as _tz
    DISPLAY_TZ = _tz(timedelta(hours=5, minutes=30))

ACTIVITY = {
    "draft.created": "created the draft",
    "check.completed": "ran the checks",
    "finding.accept": "accepted a fix",
    "finding.edit": "edited a fix",
    "finding.reject": "rejected a finding",
    "finding.acknowledge": "will fix it in the design",
    "finding.superseded": "resolved a finding with another fix",
    "version.created": "saved a new version",
    "working.discarded": "discarded working changes",
    "draft.submitted": "submitted for review",
    "draft.auto_approved": "submitted; approved automatically by policy",
    "review.approved": "approved",
    "review.changes_requested": "requested changes",
    "comment.added": "commented",
    "kit.seeded": "loaded the demo Brand Kit",
}


def _when(iso: str | None) -> str:
    if not iso:
        return ""
    from datetime import datetime
    try:
        return datetime.fromisoformat(iso).astimezone(DISPLAY_TZ).strftime("%d %b %Y, %H:%M")
    except ValueError:
        return iso


def _activity(action: str) -> str:
    if action in ACTIVITY:
        return ACTIVITY[action]
    if action.startswith("kit."):
        parts = action.split(".")
        return f"{parts[-1]} a Brand Kit {parts[1].rstrip('s')} entry" if len(parts) == 3 else action
    return action.replace(".", " ")


templates.env.filters["when"] = _when
templates.env.filters["activity"] = _activity


# --------------------------------------------------------------------------- setup
def init(database: Database | None = None) -> None:
    global db, ai
    db = database or Database()
    db.init_schema()
    seed_kit(db)
    ai = AIGateway(db)
    log.info("Brandproof ready: storage=%s, AI=%s", db.kind, ai.label)


@asynccontextmanager
async def lifespan(app):
    if db is None:
        init()
    yield


# --------------------------------------------------------------------------- helpers
def role_of(request: Request) -> str:
    r = request.cookies.get("bp_role", "creator")
    return r if r in ROLES else "creator"


def actor(request: Request) -> tuple[str, str]:
    r = role_of(request)
    return ROLES[r]["name"], r


def need_role(request: Request, *allowed: str) -> None:
    if role_of(request) not in allowed:
        names = " or ".join(ROLES[a]["label"] for a in allowed)
        raise HTTPException(403, f"Switch to {names} (top right) to do this.")


def render(request: Request, name: str, ctx: dict | None = None, status: int = 200):
    role = role_of(request)
    base = {
        "role": role, "roles": ROLES, "persona": ROLES[role], "ai_label": ai.label, "ai_on": ai.enabled,
        "storage": db.kind, "storage_warning": db.warning, "review_count": db.scalar(
            "SELECT COUNT(*) AS c FROM drafts WHERE status = 'in_review'") or 0,
        "STATUS": STATUS, "DIMENSION_LABELS": DIMENSION_LABELS, "DIMENSIONS": DIMENSIONS,
        "path": request.url.path, "flash": request.query_params.get("msg", ""),
        "flash_kind": request.query_params.get("kind", "ok"),
    }
    base.update(ctx or {})
    return templates.TemplateResponse(request, name, base, status_code=status)


def go(url: str, msg: str = "", kind: str = "ok") -> RedirectResponse:
    if msg:
        sep = "&" if "?" in url else "?"
        frag = ""
        if "#" in url:
            url, frag = url.split("#", 1)
            frag = "#" + frag
        from urllib.parse import quote
        url = f"{url}{sep}msg={quote(msg)}&kind={kind}{frag}"
    return RedirectResponse(url, status_code=303)


def get_draft(draft_id: int) -> dict:
    d = db.one("SELECT * FROM drafts WHERE id = ?", (draft_id,))
    if not d:
        raise HTTPException(404, "Draft not found")
    d["locked_list"] = json.loads(d.get("locked") or "[]")
    return d


def get_version(draft_id: int, n: int) -> dict | None:
    v = db.one("SELECT * FROM versions WHERE draft_id = ? AND n = ?", (draft_id, n))
    if v:
        v["findings_list"] = json.loads(v.get("findings") or "[]")
        v["summary_obj"] = json.loads(v.get("summary") or "{}")
    return v


def image_bytes(v: dict) -> bytes | None:
    return base64.b64decode(v["image"]) if v and v.get("image") else None


def decisions_for(draft_id: int, n: int) -> dict[str, dict]:
    rows = db.all("SELECT * FROM decisions WHERE draft_id = ? AND version_n = ? ORDER BY id", (draft_id, n))
    return {r["finding_key"]: r for r in rows}


async def read_form(request: Request):
    try:
        return await request.form(max_files=3, max_fields=80, max_part_size=MAX_UPLOAD)
    except Exception as exc:
        raise HTTPException(400, f"Could not read the form: {exc}")


async def read_upload(upload) -> bytes | None:
    if upload is None or not getattr(upload, "filename", ""):
        return None
    data = await upload.read()
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "File too large (6 MB max).")
    return data or None


def docx_text(data: bytes) -> str:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise HTTPException(400, "That file doesn't look like a valid .docx document.")
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(p for p in parts).strip()


def parse_locked(raw: str) -> list[str]:
    items = [x.strip() for x in re.split(r"[\n,]", raw or "") if x.strip()]
    return items[:10]


def run_and_store(draft: dict, n: int, body: str, image: bytes | None, image_mime: str | None, alt: str,
                  source_name: str, who: tuple[str, str]) -> dict:
    """Run every check and save the result as version n of the draft."""
    kit = load_kit(db)
    meta = {"channel": draft.get("channel"), "goal": draft.get("goal"), "locked": draft["locked_list"]}
    result = run_checks(ai, kit, meta, body, image, image_mime, alt)
    db.execute(
        """INSERT INTO versions (draft_id, n, body, image, image_mime, alt_text, source_name, findings, summary,
           kit_hash, created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (draft["id"], n, body, base64.b64encode(image).decode() if image else None, image_mime, alt, source_name,
         json.dumps(result["findings"], ensure_ascii=False), json.dumps(result["summary"], ensure_ascii=False),
         kit.hash, who[0], now_iso()),
    )
    db.execute(
        """UPDATE drafts SET current_version = ?, working_body = ?, working_alt = ?, working_dirty = 0,
           status = 'checked', updated_at = ? WHERE id = ?""",
        (n, body, alt, now_iso(), draft["id"]),
    )
    s = result["summary"]
    db.audit("check.completed", who[0], who[1], draft["id"], {
        "version": n, "issues": s["issues"], "blockers": s["counts"]["blocker"], "ai": s["ai"]["label"],
        "ai_calls_ok": sum(1 for c in s["ai"]["calls"] if c.get("ok")), "ms": s["ms"]})
    return result


def carry_rejections(draft_id: int, old_n: int, new_n: int, findings: list[dict]) -> int:
    """Keep a creator's 'reject with reason' decisions when the same issue shows up again after a re-check."""
    old = db.all("SELECT * FROM decisions WHERE draft_id = ? AND version_n = ? AND action = 'reject'", (draft_id, old_n))
    carried = 0
    for f in findings:
        for d in old:
            if d["rule_id"] == f["rule_id"] and (d.get("before_text") or "").lower() == (f.get("quote") or f["title"]).lower():
                db.execute(
                    """INSERT INTO decisions (draft_id, version_n, finding_key, rule_id, severity, action, reason,
                       before_text, after_text, actor, created_at) VALUES (?, ?, ?, ?, ?, 'reject', ?, ?, '', ?, ?)""",
                    (draft_id, new_n, f["key"], f["rule_id"], f["severity"], d["reason"], d["before_text"], d["actor"], now_iso()),
                )
                carried += 1
                break
    return carried


def highlight(body: str, findings: list[dict]) -> Markup:
    marks = [f for f in findings if f["target"] == "text" and f.get("start") is not None and f.get("end")]
    if not body:
        return Markup("")
    points = sorted({0, len(body), *[m["start"] for m in marks], *[m["end"] for m in marks]})
    rank = {"blocker": 0, "major": 1, "minor": 2, "info": 3, "ok": 4}
    out = []
    for a, b in zip(points, points[1:]):
        if a >= b:
            continue
        seg = str(escape(body[a:b])).replace("\n", "<br>")
        cover = [m for m in marks if m["start"] <= a and m["end"] >= b]
        if not cover:
            out.append(seg)
            continue
        top = min(cover, key=lambda m: rank.get(m["severity"], 9))
        keys = " ".join(m["key"] for m in cover)
        nums = [str(m["n"]) for m in cover if m["end"] == b and m.get("n")]
        sup = f"<sup>{','.join(nums)}</sup>" if nums else ""
        out.append(f'<mark class="hl hl-{top["severity"]}" data-keys="{keys}">{seg}{sup}</mark>')
    return Markup("".join(out))


def apply_fix(text: str, fix: dict, replacement: str | None, locked: list[str]) -> tuple[str | None, str]:
    """Apply a fix to the working copy. Returns (new text or None, message)."""
    kind = fix.get("type")
    if kind == "append":
        addition = (replacement if replacement is not None else fix.get("text", "")).strip()
        if addition and addition.lower() in text.lower():
            return text, "Already there"
        return (text.rstrip() + ("\n" if "\n" in text else " ") + addition).strip(), "Added"
    if kind in ("replace", "remove"):
        find = fix.get("find") or ""
        span = find_span(text, find)
        if not span:
            return None, "That text has already changed in your working copy."
        new = replacement if replacement is not None else fix.get("text", "")
        result = text[:span[0]] + new + text[span[1]:]
        if kind == "remove" or not new.strip():
            result = tidy_spaces(result)
        for phrase in locked:
            if phrase and phrase.lower() in text.lower() and phrase.lower() not in result.lower():
                return None, f"This change would remove your locked phrase “{phrase}”. Use Edit to keep it."
        return result, "Applied"
    return text, "Noted"


# --------------------------------------------------------------------------- pages
async def home(request: Request):
    drafts = db.all("SELECT * FROM drafts ORDER BY updated_at DESC, id DESC LIMIT 100")
    for d in drafts:
        v = get_version(d["id"], d["current_version"])
        d["counts"] = (v["summary_obj"].get("counts") if v else {}) or {}
        d["issues"] = (v["summary_obj"].get("issues") if v else 0) or 0
    return render(request, "index.html", {"drafts": drafts})


async def set_role(request: Request):
    form = await read_form(request)
    r = form.get("role", "creator")
    back = form.get("back") or "/"
    if not back.startswith("/"):
        back = "/"
    resp = go(back, f"You are now {ROLES.get(r, ROLES['creator'])['name']} ({ROLES.get(r, ROLES['creator'])['label']}).")
    resp.set_cookie("bp_role", r if r in ROLES else "creator", max_age=60 * 60 * 24 * 30, samesite="lax")
    return resp


async def new_draft(request: Request):
    return render(request, "new.html", {"channels": CHANNELS})


async def create_draft(request: Request):
    need_role(request, "creator", "brand")
    form = await read_form(request)
    who = actor(request)
    title = (form.get("title") or "").strip()[:140] or "Untitled draft"
    body = (form.get("body") or "").replace("\r\n", "\n").strip()
    source_name = ""
    docx_data = await read_upload(form.get("docx"))
    if docx_data:
        body = docx_text(docx_data)
        source_name = form.get("docx").filename
    img_data = await read_upload(form.get("image"))
    image, mime = None, None
    if img_data:
        try:
            image, mime, _ = normalise_image(img_data)
        except Exception:
            return go("/drafts/new", "That image could not be read. Use PNG, JPG or WebP.", "error")
    if not body and not image:
        return go("/drafts/new", "Add some text, a .docx file or an image to check.", "error")
    body = body[:MAX_TEXT]
    alt = (form.get("alt") or "").strip()[:500]
    locked = parse_locked(form.get("locked") or "")
    ts = now_iso()
    draft_id = db.insert(
        """INSERT INTO drafts (title, channel, goal, locked, status, created_by, created_at, updated_at,
           current_version, working_body, working_alt, working_dirty) VALUES (?, ?, ?, ?, 'checked', ?, ?, ?, 1, ?, ?, 0)
           RETURNING id""",
        (title, form.get("channel") or "Other", (form.get("goal") or "").strip()[:300], json.dumps(locked), who[0], ts, ts,
         body, alt),
    )
    db.audit("draft.created", who[0], who[1], draft_id, {"title": title, "docx": bool(docx_data), "image": bool(image)})
    draft = get_draft(draft_id)
    await run_in_threadpool(run_and_store, draft, 1, body, image, mime, alt, source_name, who)
    return go(f"/drafts/{draft_id}")


async def create_sample(request: Request):
    need_role(request, "creator", "brand")
    name = request.path_params["name"]
    sample = SAMPLE_POST if name == "instagram" else SAMPLE_EMAIL if name == "email" else None
    if not sample:
        raise HTTPException(404, "Unknown sample")
    who = actor(request)
    image, mime = None, None
    if sample.get("image"):
        with open(os.path.join(HERE, "static", sample["image"]), "rb") as fh:
            image, mime, _ = normalise_image(fh.read())
    ts = now_iso()
    draft_id = db.insert(
        """INSERT INTO drafts (title, channel, goal, locked, status, created_by, created_at, updated_at,
           current_version, working_body, working_alt, working_dirty) VALUES (?, ?, ?, ?, 'checked', ?, ?, ?, 1, ?, ?, 0)
           RETURNING id""",
        (sample["title"], sample["channel"], sample["goal"], json.dumps(sample["locked"]), who[0], ts, ts,
         sample["body"], sample["alt"]),
    )
    db.audit("draft.created", who[0], who[1], draft_id, {"title": sample["title"], "sample": name})
    draft = get_draft(draft_id)
    await run_in_threadpool(run_and_store, draft, 1, sample["body"], image, mime, sample["alt"], "", who)
    return go(f"/drafts/{draft_id}")


async def draft_page(request: Request):
    draft_id = int(request.path_params["id"])
    draft = get_draft(draft_id)
    n = int(request.query_params.get("v") or draft["current_version"])
    n = max(1, min(n, draft["current_version"]))
    version = get_version(draft_id, n)
    if not version:
        raise HTTPException(404, "Version not found")
    is_current = n == draft["current_version"]
    findings = version["findings_list"]
    decisions = decisions_for(draft_id, n)
    issues = [f for f in findings if f["severity"] != "ok"]
    supported = [f for f in findings if f["severity"] == "ok"]
    open_blockers = [f for f in issues if f["severity"] == "blocker" and f["key"] not in decisions]
    versions = db.all("SELECT n, created_by, created_at, summary FROM versions WHERE draft_id = ? ORDER BY n DESC", (draft_id,))
    for v in versions:
        s = json.loads(v["summary"] or "{}")
        v["issues"] = s.get("issues", 0)
        v["blockers"] = (s.get("counts") or {}).get("blocker", 0)
    prev = get_version(draft_id, n - 1) if n > 1 else None
    diff = Markup(word_diff_html(prev["body"] or "", version["body"] or "")) if prev else None
    working_diff = None
    if is_current and draft["working_dirty"]:
        working_diff = Markup(word_diff_html(version["body"] or "", draft["working_body"] or ""))
    comments = db.all("SELECT * FROM comments WHERE draft_id = ? ORDER BY id", (draft_id,))
    audit = db.all("SELECT * FROM audit WHERE draft_id = ? ORDER BY id DESC LIMIT 60", (draft_id,))
    for a in audit:
        a["detail_obj"] = json.loads(a.get("detail") or "{}")
    approval = None
    if draft["status"] == "approved":
        approval = db.one("SELECT * FROM audit WHERE draft_id = ? AND action IN ('review.approved', 'draft.auto_approved') "
                          "ORDER BY id DESC", (draft_id,))
        if approval:
            approval["detail_obj"] = json.loads(approval.get("detail") or "{}")
    role = role_of(request)
    ctx = {
        "draft": draft, "version": version, "n": n, "is_current": is_current, "findings": issues,
        "supported": supported, "decisions": decisions, "open_blockers": open_blockers, "versions": versions,
        "diff": diff, "working_diff": working_diff, "comments": comments, "audit": audit, "approval": approval,
        "body_html": highlight(version["body"] or "", findings),
        "can_edit": is_current and draft["status"] in EDITABLE and role in ("creator", "brand"),
        "can_review": is_current and draft["status"] == "in_review" and role in ("reviewer", "brand"),
        "summary": version["summary_obj"], "channels": CHANNELS,
    }
    return render(request, "draft.html", ctx)


async def version_image(request: Request):
    v = get_version(int(request.path_params["id"]), int(request.path_params["n"]))
    data = image_bytes(v)
    if not data:
        raise HTTPException(404, "No image")
    return Response(data, media_type=v.get("image_mime") or "image/png", headers={"Cache-Control": "public, max-age=86400"})


async def decide(request: Request):
    need_role(request, "creator", "brand")
    draft_id = int(request.path_params["id"])
    key = request.path_params["key"]
    draft = get_draft(draft_id)
    if draft["status"] not in EDITABLE:
        return go(f"/drafts/{draft_id}", "This draft is locked while it is in review or approved.", "error")
    version = get_version(draft_id, draft["current_version"])
    finding = next((f for f in version["findings_list"] if f["key"] == key), None)
    if not finding:
        raise HTTPException(404, "Finding not found")
    form = await read_form(request)
    action = form.get("action")
    who = actor(request)
    anchor = f"/drafts/{draft_id}#f-{key}"
    if decisions_for(draft_id, draft["current_version"]).get(key):
        return go(anchor, "You already decided on this finding.", "error")
    before = finding.get("quote") or finding["title"]
    after = ""
    fix = finding.get("fix") or {"type": "advice", "text": ""}
    if action == "reject":
        reason = (form.get("reason") or "").strip()
        if len(reason) < 3:
            return go(anchor, "Please give a short reason for rejecting this finding.", "error")
        msg = "Rejected. Your reason is saved and shown to the reviewer."
    elif action in ("accept", "edit"):
        reason = ""
        replacement = form.get("text") if action == "edit" else None
        if replacement is not None:
            replacement = replacement.strip()
        if fix["type"] == "alt":
            new_alt = replacement if replacement is not None else fix.get("text", "")
            db.execute("UPDATE drafts SET working_alt = ?, working_dirty = 1, updated_at = ? WHERE id = ?",
                       (new_alt, now_iso(), draft_id))
            after = new_alt
            msg = "Alt text added to your working copy. Re-check to confirm."
        elif fix["type"] in ("replace", "remove", "append"):
            new_text, msg = apply_fix(draft["working_body"] or "", fix, replacement, draft["locked_list"])
            if new_text is None:
                return go(anchor, msg, "error")
            after = replacement if replacement is not None else fix.get("text", "")
            db.execute("UPDATE drafts SET working_body = ?, working_dirty = 1, updated_at = ? WHERE id = ?",
                       (new_text, now_iso(), draft_id))
            msg = f"{msg} to your working copy. Re-check when you're ready."
        else:
            action = "acknowledge"
            msg = "Noted. Fix it in the design, then upload the new image when you re-check."
    else:
        raise HTTPException(400, "Unknown action")
    db.execute(
        """INSERT INTO decisions (draft_id, version_n, finding_key, rule_id, severity, action, reason, before_text,
           after_text, actor, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (draft_id, draft["current_version"], key, finding["rule_id"], finding["severity"], action, reason, before, after,
         who[0], now_iso()),
    )
    db.audit(f"finding.{action}", who[0], who[1], draft_id, {
        "version": draft["current_version"], "rule": finding["rule_id"], "finding": finding["title"],
        "reason": reason, "after": after[:200]})
    if action in ("accept", "edit") and fix["type"] in ("replace", "remove") and finding.get("start") is not None:
        # findings inside the text that was just replaced are resolved by this fix
        decided = decisions_for(draft_id, draft["current_version"])
        superseded = 0
        for other in version["findings_list"]:
            if (other["key"] == key or other["key"] in decided or other.get("target") != "text"
                    or other.get("start") is None or other["severity"] == "ok"
                    or (other.get("fix") or {}).get("type") not in ("replace", "remove")):
                continue
            if finding["start"] <= other["start"] and other["end"] <= finding["end"]:
                db.execute(
                    """INSERT INTO decisions (draft_id, version_n, finding_key, rule_id, severity, action, reason,
                       before_text, after_text, actor, created_at) VALUES (?, ?, ?, ?, ?, 'superseded', ?, ?, '', ?, ?)""",
                    (draft_id, draft["current_version"], other["key"], other["rule_id"], other["severity"],
                     f"Resolved by the fix for finding {finding.get('n') or ''}", other.get("quote") or other["title"],
                     who[0], now_iso()),
                )
                superseded += 1
        if superseded:
            msg += f" {superseded} overlapping finding(s) were resolved by the same fix."
    return go(anchor, msg)


async def discard_working(request: Request):
    need_role(request, "creator", "brand")
    draft_id = int(request.path_params["id"])
    draft = get_draft(draft_id)
    version = get_version(draft_id, draft["current_version"])
    db.execute("UPDATE drafts SET working_body = ?, working_alt = ?, working_dirty = 0 WHERE id = ?",
               (version["body"], version["alt_text"], draft_id))
    db.execute("DELETE FROM decisions WHERE draft_id = ? AND version_n = ? AND action IN ('accept', 'edit')",
               (draft_id, draft["current_version"]))
    who = actor(request)
    db.audit("working.discarded", who[0], who[1], draft_id, {"version": draft["current_version"]})
    return go(f"/drafts/{draft_id}", "Working copy reset to the last checked version.")


async def recheck(request: Request):
    need_role(request, "creator", "brand")
    draft_id = int(request.path_params["id"])
    draft = get_draft(draft_id)
    if draft["status"] not in EDITABLE:
        return go(f"/drafts/{draft_id}", "This draft is locked while it is in review or approved.", "error")
    form = await read_form(request)
    who = actor(request)
    old = get_version(draft_id, draft["current_version"])
    body = form.get("body")
    body = (body if body is not None else draft["working_body"] or "").replace("\r\n", "\n").strip()[:MAX_TEXT]
    alt = form.get("alt")
    alt = (alt if alt is not None else draft["working_alt"] or "").strip()[:500]
    image, mime = image_bytes(old), old.get("image_mime")
    new_img = await read_upload(form.get("image"))
    if new_img:
        try:
            image, mime, _ = normalise_image(new_img)
        except Exception:
            return go(f"/drafts/{draft_id}", "That image could not be read. Use PNG, JPG or WebP.", "error")
    if form.get("remove_image") == "1":
        image, mime = None, None
    if not body and not image:
        return go(f"/drafts/{draft_id}", "The draft is empty.", "error")
    n = draft["current_version"] + 1
    db.audit("version.created", who[0], who[1], draft_id, {"version": n, "new_image": bool(new_img)})
    result = await run_in_threadpool(run_and_store, draft, n, body, image, mime, alt, old.get("source_name") or "", who)
    carried = carry_rejections(draft_id, n - 1, n, result["findings"])
    s = result["summary"]
    msg = f"Version {n} checked: {s['issues']} issue(s) left."
    if carried:
        msg += f" {carried} earlier rejection(s) carried over."
    return go(f"/drafts/{draft_id}", msg)


def routing_decision(version: dict, decisions: dict[str, dict]) -> tuple[bool, list[str]]:
    reasons = []
    findings = version["findings_list"]
    if any(f["dimension"] == "claims" for f in findings):
        reasons.append("it makes factual or promotional claims")
    overridden = [f for f in findings if decisions.get(f["key"], {}).get("action") == "reject"
                  and f["severity"] in ("blocker", "major")]
    if overridden:
        reasons.append(f"{len(overridden)} serious finding(s) were rejected by the creator")
    if any(f.get("needs_human") for f in findings):
        reasons.append("the AI was unsure about something and asked for human judgment")
    open_major = [f for f in findings if f["severity"] == "major" and f["key"] not in decisions]
    if open_major:
        reasons.append(f"{len(open_major)} major finding(s) are still open")
    return bool(reasons), reasons


async def submit(request: Request):
    need_role(request, "creator", "brand")
    draft_id = int(request.path_params["id"])
    draft = get_draft(draft_id)
    who = actor(request)
    if draft["status"] not in EDITABLE:
        return go(f"/drafts/{draft_id}", "This draft has already been submitted.", "error")
    if draft["working_dirty"]:
        return go(f"/drafts/{draft_id}", "You have unchecked changes. Re-check before submitting.", "error")
    version = get_version(draft_id, draft["current_version"])
    decisions = decisions_for(draft_id, draft["current_version"])
    blockers = [f for f in version["findings_list"] if f["severity"] == "blocker" and f["key"] not in decisions]
    if blockers:
        return go(f"/drafts/{draft_id}", f"Resolve {len(blockers)} blocker(s) first: accept a fix and re-check, "
                  "or reject with a reason.", "error")
    needs, reasons = routing_decision(version, decisions)
    if needs:
        reason_text = "; ".join(reasons)
        db.execute("UPDATE drafts SET status = 'in_review', review_reason = ?, updated_at = ? WHERE id = ?",
                   (reason_text, now_iso(), draft_id))
        db.audit("draft.submitted", who[0], who[1], draft_id, {"version": draft["current_version"], "routed_because": reasons})
        return go(f"/drafts/{draft_id}", f"Sent to a reviewer because {reason_text}.")
    db.execute("UPDATE drafts SET status = 'approved', review_reason = ?, updated_at = ? WHERE id = ?",
               ("Auto-cleared: no claims, no open serious issues", now_iso(), draft_id))
    db.audit("draft.auto_approved", who[0], who[1], draft_id, {
        "version": draft["current_version"], "policy": "No claims and no open blocker or major findings"})
    return go(f"/drafts/{draft_id}", "Approved automatically: no claims and no serious issues, so no review was needed.")


async def review(request: Request):
    need_role(request, "reviewer", "brand")
    draft_id = int(request.path_params["id"])
    draft = get_draft(draft_id)
    if draft["status"] != "in_review":
        return go(f"/drafts/{draft_id}", "This draft is not waiting for review.", "error")
    form = await read_form(request)
    who = actor(request)
    note = (form.get("comment") or "").strip()[:2000]
    if note:
        db.execute("INSERT INTO comments (draft_id, actor, role, body, created_at) VALUES (?, ?, ?, ?, ?)",
                   (draft_id, who[0], who[1], note, now_iso()))
    if form.get("decision") == "approve":
        version = get_version(draft_id, draft["current_version"])
        decisions = decisions_for(draft_id, draft["current_version"])
        overrides = [{"rule": f["rule_id"], "finding": f["title"], "reason": decisions[f["key"]]["reason"]}
                     for f in version["findings_list"] if decisions.get(f["key"], {}).get("action") == "reject"]
        db.execute("UPDATE drafts SET status = 'approved', updated_at = ? WHERE id = ?", (now_iso(), draft_id))
        db.audit("review.approved", who[0], who[1], draft_id, {
            "version": draft["current_version"], "comment": note, "overrides_accepted": overrides,
            "ai": version["summary_obj"].get("ai", {}).get("label", ""), "kit": version.get("kit_hash")})
        return go(f"/drafts/{draft_id}", f"Approved version {draft['current_version']}. It is now locked.")
    if form.get("decision") == "changes":
        if not note:
            return go(f"/drafts/{draft_id}", "Tell the creator what to change (comment required).", "error")
        db.execute("UPDATE drafts SET status = 'changes_requested', updated_at = ? WHERE id = ?", (now_iso(), draft_id))
        db.audit("review.changes_requested", who[0], who[1], draft_id, {"version": draft["current_version"], "comment": note})
        return go(f"/drafts/{draft_id}", "Changes requested. The creator can now edit and resubmit.")
    raise HTTPException(400, "Unknown decision")


async def add_comment(request: Request):
    draft_id = int(request.path_params["id"])
    get_draft(draft_id)
    form = await read_form(request)
    body = (form.get("body") or "").strip()[:2000]
    if not body:
        return go(f"/drafts/{draft_id}#comments", "Write a comment first.", "error")
    who = actor(request)
    db.execute("INSERT INTO comments (draft_id, actor, role, body, created_at) VALUES (?, ?, ?, ?, ?)",
               (draft_id, who[0], who[1], body, now_iso()))
    db.audit("comment.added", who[0], who[1], draft_id, {"comment": body[:200]})
    return go(f"/drafts/{draft_id}#comments", "Comment added.")


async def review_queue(request: Request):
    drafts = db.all("SELECT * FROM drafts WHERE status = 'in_review' ORDER BY updated_at")
    recent = db.all("SELECT * FROM drafts WHERE status IN ('approved', 'changes_requested') ORDER BY updated_at DESC LIMIT 10")
    return render(request, "review.html", {"drafts": drafts, "recent": recent})


# --------------------------------------------------------------------------- brand kit
KIT_TABLES = {
    "rules": ["id", "dimension", "title", "description", "severity", "method", "pattern", "fix_hint", "do_example",
              "dont_example", "source", "active"],
    "glossary": ["preferred", "avoid", "note", "rule_id"],
    "facts": ["id", "topic", "label", "statement", "value", "approved_wording", "keywords", "evidence", "owner",
              "valid_until"],
    "palette": ["name", "hex"],
}


async def kit_page(request: Request):
    tab = request.query_params.get("tab", "rules")
    if tab not in KIT_TABLES:
        tab = "rules"
    kit = load_kit(db)
    edit_id = request.query_params.get("edit")
    rows = {
        "rules": db.all("SELECT * FROM rules ORDER BY id"),
        "glossary": db.all("SELECT * FROM glossary ORDER BY id"),
        "facts": db.all("SELECT * FROM facts ORDER BY id"),
        "palette": db.all("SELECT * FROM palette ORDER BY id"),
    }
    today = date.today()
    for f in rows["facts"]:
        try:
            days = (date.fromisoformat(f["valid_until"]) - today).days if f.get("valid_until") else None
        except ValueError:
            days = None
        f["days_left"] = days
    editing = None
    if edit_id:
        key = "id"
        editing = next((r for r in rows[tab] if str(r.get(key)) == edit_id), None)
    return render(request, "kit.html", {"tab": tab, "rows": rows, "editing": editing, "kit_hash": kit.hash,
                                        "fields": KIT_TABLES, "disclaimer": kit.disclaimer})


def _clean_row(table: str, form) -> dict:
    row = {k: (form.get(k) or "").strip() for k in KIT_TABLES[table]}
    if table == "rules":
        if not re.fullmatch(r"[A-Z]-[A-Z]{3}-\d{2,3}|R-[A-Z0-9-]{3,20}", row["id"] or ""):
            raise ValueError("Rule ID must look like R-TON-05.")
        if row["dimension"] not in DIMENSIONS:
            raise ValueError("Pick a valid dimension.")
        if row["severity"] not in ("blocker", "major", "minor"):
            raise ValueError("Severity must be blocker, major or minor.")
        if row["method"] not in ("code", "ai", "code+ai"):
            raise ValueError("Method must be code, ai or code+ai.")
        if row["pattern"]:
            try:
                re.compile(row["pattern"])
            except re.error as exc:
                raise ValueError(f"Pattern is not a valid regular expression: {exc}")
        if not row["title"]:
            raise ValueError("Give the rule a title.")
        raw_active = form.get("active")
        if form.get("active_present"):
            row["active"] = 1 if raw_active else 0
        else:  # CSV import: missing or blank means active
            row["active"] = 0 if str(raw_active or "1").strip().lower() in ("0", "false", "no", "off") else 1
    elif table == "facts":
        if not re.fullmatch(r"F-[A-Za-z0-9-]{1,20}", row["id"] or ""):
            raise ValueError("Fact ID must look like F-14.")
        if not row["label"] or not row["statement"]:
            raise ValueError("A fact needs a label and a statement.")
        if row["valid_until"]:
            date.fromisoformat(row["valid_until"])
        if row["value"]:
            float(row["value"].replace(",", ""))
    elif table == "palette":
        if not re.fullmatch(r"#?[0-9A-Fa-f]{6}", row["hex"]):
            raise ValueError("Colour must be a hex code like #11734F.")
        row["hex"] = "#" + row["hex"].lstrip("#").upper()
        if not row["name"]:
            raise ValueError("Give the colour a name.")
    elif table == "glossary":
        if not row["preferred"] or not row["avoid"]:
            raise ValueError("A glossary entry needs a preferred term and words to avoid.")
        row["rule_id"] = row["rule_id"] or "R-TRM-01"
    return row


def _save_row(table: str, row: dict, row_id: str | None = None) -> None:
    ts = now_iso()
    cols = KIT_TABLES[table]
    if table in ("rules", "facts"):
        sets = ", ".join(f"{c} = excluded.{c}" for c in cols if c != "id")
        db.execute(
            f"INSERT INTO {table} ({', '.join(cols)}, updated_at) VALUES ({', '.join('?' for _ in cols)}, ?) "
            f"ON CONFLICT (id) DO UPDATE SET {sets}, updated_at = excluded.updated_at",
            (*[row[c] for c in cols], ts),
        )
    elif row_id:
        db.execute(f"UPDATE {table} SET {', '.join(f'{c} = ?' for c in cols)}, updated_at = ? WHERE id = ?",
                   (*[row[c] for c in cols], ts, int(row_id)))
    else:
        db.execute(f"INSERT INTO {table} ({', '.join(cols)}, updated_at) VALUES ({', '.join('?' for _ in cols)}, ?)",
                   (*[row[c] for c in cols], ts))


async def kit_save(request: Request):
    need_role(request, "brand")
    table = request.path_params["table"]
    if table not in KIT_TABLES:
        raise HTTPException(404, "Unknown table")
    form = await read_form(request)
    who = actor(request)
    try:
        row = _clean_row(table, form)
        _save_row(table, row, form.get("row_id") or None)
    except (ValueError, TypeError) as exc:
        return go(f"/kit?tab={table}", str(exc), "error")
    db.audit(f"kit.{table}.saved", who[0], who[1], None, {"row": {k: str(v)[:80] for k, v in row.items()}})
    return go(f"/kit?tab={table}", "Saved. New checks use it straight away.")


async def kit_delete(request: Request):
    need_role(request, "brand")
    table = request.path_params["table"]
    if table not in KIT_TABLES:
        raise HTTPException(404, "Unknown table")
    form = await read_form(request)
    row_id = form.get("row_id")
    who = actor(request)
    if table in ("rules", "facts"):
        db.execute(f"DELETE FROM {table} WHERE id = ?", (row_id,))
    else:
        db.execute(f"DELETE FROM {table} WHERE id = ?", (int(row_id),))
    db.audit(f"kit.{table}.deleted", who[0], who[1], None, {"id": row_id})
    return go(f"/kit?tab={table}", "Deleted.")


async def kit_export(request: Request):
    table = request.path_params["table"]
    if table not in KIT_TABLES:
        raise HTTPException(404, "Unknown table")
    rows = db.all(f"SELECT {', '.join(KIT_TABLES[table])} FROM {table} ORDER BY 1")
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=KIT_TABLES[table])
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return PlainTextResponse(buf.getvalue(), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="brandproof-{table}.csv"'})


async def kit_import(request: Request):
    need_role(request, "brand")
    table = request.path_params["table"]
    if table not in KIT_TABLES:
        raise HTTPException(404, "Unknown table")
    form = await read_form(request)
    data = await read_upload(form.get("file"))
    if not data:
        return go(f"/kit?tab={table}", "Choose a CSV file first.", "error")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return go(f"/kit?tab={table}", "The CSV must be UTF-8 encoded.", "error")
    reader = csv.DictReader(io.StringIO(text))
    missing = [c for c in KIT_TABLES[table] if c not in (reader.fieldnames or []) and c not in ("active",)]
    if missing:
        return go(f"/kit?tab={table}", f"CSV is missing columns: {', '.join(missing)}", "error")
    ok, errors = 0, []
    for i, raw in enumerate(reader, start=2):
        try:
            row = _clean_row(table, raw)
            _save_row(table, row)
            ok += 1
        except (ValueError, TypeError) as exc:
            errors.append(f"line {i}: {exc}")
    who = actor(request)
    db.audit(f"kit.{table}.imported", who[0], who[1], None, {"rows": ok, "errors": errors[:5]})
    msg = f"Imported {ok} row(s)." + (f" Skipped {len(errors)}: {'; '.join(errors[:3])}" if errors else "")
    return go(f"/kit?tab={table}", msg, "error" if errors and not ok else "ok")


# --------------------------------------------------------------------------- other pages
async def audit_page(request: Request):
    rows = db.all("SELECT a.*, d.title AS draft_title FROM audit a LEFT JOIN drafts d ON d.id = a.draft_id "
                  "ORDER BY a.id DESC LIMIT 300")
    for r in rows:
        r["detail_obj"] = json.loads(r.get("detail") or "{}")
    return render(request, "audit.html", {"rows": rows})


async def about_page(request: Request):
    calls = db.one("SELECT COUNT(*) AS n, SUM(ok) AS ok, SUM(cached) AS cached, AVG(ms) AS ms FROM ai_calls") or {}
    recent_errors = db.all("SELECT purpose, error, created_at FROM ai_calls WHERE ok = 0 ORDER BY id DESC LIMIT 5")
    accuracy = None
    path = os.path.join(os.path.dirname(HERE), "data", "eval_results.json")
    if os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            accuracy = json.load(fh)
    return render(request, "about.html", {"calls": calls, "recent_errors": recent_errors, "accuracy": accuracy,
                                          "provider": ai.provider, "model": ai.model})


async def healthz(request: Request):
    return JSONResponse({"ok": True, "storage": db.kind, "ai": ai.label, "warning": db.warning})


async def http_error(request: Request, exc: HTTPException):
    return render(request, "error.html", {"code": exc.status_code, "detail": exc.detail}, status=exc.status_code)


async def server_error(request: Request, exc: Exception):
    log.exception("Unhandled error on %s", request.url.path)
    return render(request, "error.html", {"code": 500, "detail": "Something went wrong on our side. Please try again."},
                  status=500)


routes = [
    Route("/", home),
    Route("/role", set_role, methods=["POST"]),
    Route("/drafts/new", new_draft),
    Route("/drafts", create_draft, methods=["POST"]),
    Route("/samples/{name}", create_sample, methods=["POST"]),
    Route("/drafts/{id:int}", draft_page),
    Route("/drafts/{id:int}/v/{n:int}/image", version_image),
    Route("/drafts/{id:int}/findings/{key}", decide, methods=["POST"]),
    Route("/drafts/{id:int}/discard", discard_working, methods=["POST"]),
    Route("/drafts/{id:int}/recheck", recheck, methods=["POST"]),
    Route("/drafts/{id:int}/submit", submit, methods=["POST"]),
    Route("/drafts/{id:int}/review", review, methods=["POST"]),
    Route("/drafts/{id:int}/comments", add_comment, methods=["POST"]),
    Route("/review", review_queue),
    Route("/kit", kit_page),
    Route("/kit/{table}/save", kit_save, methods=["POST"]),
    Route("/kit/{table}/delete", kit_delete, methods=["POST"]),
    Route("/kit/{table}/import", kit_import, methods=["POST"]),
    Route("/kit/{table}.csv", kit_export),
    Route("/audit", audit_page),
    Route("/about", about_page),
    Route("/healthz", healthz),
    Mount("/static", StaticFiles(directory=os.path.join(HERE, "static")), name="static"),
]

app = Starlette(routes=routes, lifespan=lifespan, exception_handlers={HTTPException: http_error, 500: server_error},
                max_body_size=8 * 1024 * 1024)
